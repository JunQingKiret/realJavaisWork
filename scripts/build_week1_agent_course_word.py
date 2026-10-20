from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
import re
import sys
from zipfile import ZipFile
from xml.etree import ElementTree as ET

SOURCE = Path('deliverables/第一周Agent协作与Java工具验证课程指导方案.md')
OUTPUT = Path('deliverables/第一周Agent协作与Java工具验证课程指导方案-教师执行稿.docx')
NAVY = '17365D'
GRAY = '596579'


def set_font(style, latin='Times New Roman', east='宋体', size=11, bold=False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for key, val in [('ascii', latin), ('hAnsi', latin), ('eastAsia', east)]:
        rf.set(qn('w:' + key), val)


def run_font(run, latin=None, east=None, size=None, color=None):
    if latin:
        run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if east:
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        rf.set(qn('w:eastAsia'), east)


def plain_inline(text):
    return re.sub(r'\*\*(.+?)\*\*|`([^`]+)`', lambda m: m.group(1) if m.group(1) is not None else m.group(2), text)


def add_inline(paragraph, text):
    parts = re.split(r'(\*\*.+?\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run_font(run, latin='Consolas', size=9.5, color=GRAY)
        else:
            paragraph.add_run(part)


def shade(element, fill):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    element.append(shading)


def field(paragraph, name):
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' ' + name + ' '
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    result = OxmlElement('w:t')
    result.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in [begin, instruction, sep, result, end]:
        paragraph.add_run()._r.append(el)


def table_cells(line):
    raw = line.strip().strip('|')
    return [s.strip().replace('\\|', '|') for s in re.split(r'(?<!\\)\|', raw)]


def is_separator(cells):
    return all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in cells)


def render_table(doc, lines, expected):
    rows = [table_cells(line) for line in lines]
    assert len(rows) >= 2 and is_separator(rows[1]), 'Invalid Markdown table'
    rows.pop(1)
    width = len(rows[0])
    assert all(len(r) == width for r in rows), 'Inconsistent table column count'
    table = doc.add_table(rows=1, cols=width)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    ratios = {2: [0.26, 0.74], 3: [0.23, 0.39, 0.38], 4: [0.12, 0.24, 0.34, 0.30]}.get(width, [1/width]*width)
    for col, ratio in zip(table.columns, ratios):
        col.width = Mm(168*ratio)
    for row_index, data in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        row_properties.append(no_split)
        if row_index == 0:
            repeat = OxmlElement('w:tblHeader')
            row_properties.append(repeat)
        for cell, ratio, text in zip(row.cells, ratios, data):
            cell.width = Mm(168*ratio)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for name, value in [('top', '70'), ('left', '90'), ('bottom', '70'), ('right', '90')]:
                setting = OxmlElement('w:' + name)
                setting.set(qn('w:w'), value)
                setting.set(qn('w:type'), 'dxa')
                margins.append(setting)
            tcpr.append(margins)
            if row_index == 0:
                shade(tcpr, 'DCE6F1')
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_with_next = False
            for j, part in enumerate(text.split('<br>')):
                if j:
                    p.add_run().add_break()
                add_inline(p, part)
            for run in p.runs:
                run_font(run, size=9.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
            expected.append(plain_inline(text).replace('<br>', ''))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    after.paragraph_format.space_before = Pt(0)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(22)
    section.right_margin = Mm(20)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(9)
    normal = doc.styles['Normal']
    set_font(normal, size=11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for level, size in [(1, 15), (2, 12.5), (3, 11.5)]:
        style = doc.styles[f'Heading {level}']
        set_font(style, latin='Arial', east='微软雅黑', size=size, bold=True)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run('第一周 Agent 协作与 Java 工具验证  |  教师执行稿')
    run_font(hr, east='宋体', size=8, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    field(footer, 'PAGE')
    footer.add_run(' 页 / 共 ')
    field(footer, 'NUMPAGES')
    footer.add_run(' 页')
    for r in footer.runs:
        run_font(r, east='宋体', size=8, color=GRAY)
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.append(update)
    doc.core_properties.title = '第一周 Agent 协作与 Java 工具验证课程指导方案'
    doc.core_properties.subject = '10人线上与线下结合的第一周Agent协作课程执行方案'
    doc.core_properties.author = '教学共备'
    doc.core_properties.keywords = 'Agent, Java, CodeArts, 腾讯会议, 第一周课程, 10人教学'
    doc.core_properties.comments = '由Markdown源稿转换；周二、周三和总结汇报具体时间待教师确认。'
    lines = SOURCE.read_text(encoding='utf-8-sig').splitlines()
    expected = []
    in_code = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if not line:
            i += 1
            continue
        if in_code:
            text = line
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            shade(p._p.get_or_add_pPr(), 'F2F4F7')
            r = p.add_run(text)
            run_font(r, latin='Consolas', east='微软雅黑', size=9.5, color=GRAY)
            expected.append(text)
            i += 1
            continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            render_table(doc, table_lines, expected)
            continue
        if line.startswith('# '):
            text = line[2:]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.bold = True
            run_font(r, latin='Arial', east='微软雅黑', size=20, color=NAVY)
        elif line.startswith('## '):
            text = line[3:]
            p = doc.add_heading(text, level=1)
        elif line.startswith('#### '):
            text = line[5:]
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.bold = True
            run_font(r, east='微软雅黑', size=11, color=NAVY)
        elif line.startswith('### '):
            text = line[4:]
            p = doc.add_heading(text, level=2)
        elif line.startswith('> '):
            text = line[2:]
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(8)
            shade(p._p.get_or_add_pPr(), 'F2F4F7')
            add_inline(p, text)
            for r in p.runs:
                run_font(r, size=9.5, color=GRAY)
        else:
            text = line
            p = doc.add_paragraph()
            if line.startswith('- '):
                text = line[2:]
                p.paragraph_format.left_indent = Mm(4)
                p.paragraph_format.first_line_indent = Mm(-3)
                p.add_run('• ')
            elif re.match(r'^\d+\.\s', line):
                p.paragraph_format.left_indent = Mm(5)
                p.paragraph_format.first_line_indent = Mm(-5)
            elif line.startswith('**实训步骤'):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(8)
            add_inline(p, text)
            if line.startswith('**学习主线'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
            elif not line.startswith(('- ', '**实训步骤')):
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        expected.append(plain_inline(text))
        i += 1
    doc.save(OUTPUT)
    reopened = Document(OUTPUT)
    all_text = ''.join(node.text or '' for node in reopened.element.body.iter(qn('w:t')))
    missing = [s for s in expected if s and s not in all_text]
    assert not missing, f'Markdown text lost: {missing}'
    chapter_titles = [p.text for p in reopened.paragraphs if p.style.name == 'Heading 1']
    assert len(chapter_titles) == 8, chapter_titles
    assert '一、任务目标' in chapter_titles and '八、思考题' in chapter_titles
    expected_tables = sum(line.strip().startswith('|') and (i == 0 or not lines[i-1].strip().startswith('|')) for i, line in enumerate(lines))
    assert len(reopened.tables) == expected_tables, (len(reopened.tables), expected_tables)
    with ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith('.xml'):
                ET.fromstring(archive.read(name))
    report = {
        'markdown_source': str(SOURCE),
        'word_output': str(OUTPUT),
        'chapters': chapter_titles,
        'table_count': len(reopened.tables),
        'all_markdown_text_preserved': not missing,
        'paragraph_count_outside_tables': len(reopened.paragraphs),
        'word_bytes': OUTPUT.stat().st_size,
        'paper': 'A4 portrait',
        'body_font': '宋体 11pt / Times New Roman',
        'heading_font': '微软雅黑',
        'table_font_pt': 9.5,
        'header_rows_repeat': True,
        'page_number_fields': ['PAGE', 'NUMPAGES'],
        'zip_and_xml_valid': True,
        'rendered_page_review': 'not performed; document structure and text checked',
    }
    Path('deliverables/第一周课程Word生成校验.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    build()


