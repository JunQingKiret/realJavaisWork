from pathlib import Path
from xml.sax.saxutils import escape
import json, hashlib, zipfile, re
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image, ImageOps, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
STUDENT = ROOT / 'student'
SOURCE = ROOT / 'source'
QA = ROOT / 'qa'
FONT = 'CourseCN'
pdfmetrics.registerFont(TTFont(FONT, 'C:/Windows/Fonts/simhei.ttf'))
pdfmetrics.registerFontFamily(FONT, normal=FONT, bold=FONT, italic=FONT, boldItalic=FONT)
INK, MUTED, TEAL, PALE = '#203348', '#627487', '#087E8B', '#EFF6F7'

# All sources are retained so the teacher can revise the handouts later.
def p(text): return {'type':'p', 'text':text}
def h(text): return {'type':'h', 'text':text}
def note(text): return {'type':'note', 'text':text}
def prompt(text): return {'type':'prompt', 'text':text}
def table(headers, rows, widths): return {'type':'table', 'headers':headers, 'rows':rows, 'widths':widths}
def lines(labels): return {'type':'lines', 'labels':labels}
def page(title, subtitle, *blocks): return {'title':title, 'subtitle':subtitle, 'blocks':list(blocks)}

HANDBOOK = [
page('从想法到第一版仓库系统', '第 2 课 · 线下实战 · 课前预习手册',
 note('本课目标：先理清需求，再尝试用 Agent 搭建仓库系统；逐步增加功能，在运行和验证中发现问题。'),
 h('这次，我们继续做 103 仓库'),
 p('上一课，我们用 Java 的输入、分支、循环、方法和 List，练习查看与搜索物品。这一课从“程序能查东西”继续往前走：谁来使用？需要哪些功能？怎样借助 Agent 做出第一版？'),
 h('课堂会做什么'),
 table(['环节','你将参与的事情'],[
 ['认识 Agent','比较对话式、通用型、编码型的使用场景。'],
 ['练习提问','补充提示词与上下文，让任务清楚、结果可检查。'],
 ['模拟甲乙方','提出需求、追问规则、确定本次范围。'],
 ['尝试搭建','让 Agent 分步实现，运行系统并记录问题。'],
 ['穿插学习','遇到实际问题时，再连接 Java、JavaWeb 与 MySQL 知识。']],[0.23,0.77]),
 h('课前准备约 20–30 分钟'),
 p('① 约 5 分钟：打开上一课程序，尝试运行并记录卡点。\n② 约 10–15 分钟：浏览本手册，不要求背术语。\n③ 约 5–10 分钟：填写工作单第 1 页，带着问题来上课。'),
 note('需要带：电脑与充电器、上一课项目、填写后的工作单。上课时间和地点以老师通知为准。课前不用把仓库系统做出来。'),
),
page('Agent：分清任务与执行结果', '先认识用法，再学习操作',
 p('本课把对话式 Agent、通用型 Agent、编码型 Agent 作为便于理解的使用场景。这些能力可以重叠，不是三种互斥的技术类别。'),
 table(['使用场景','仓库项目里的例子','你要检查什么'],[
 ['对话式助手','讨论需求；解释一段代码或报错。','是否结合了我们的实际情况？'],
 ['通用任务 Agent','整理访谈材料；形成需求文档和任务清单。','是否把猜测写成了已确认需求？'],
 ['编码 Agent','查看工程；修改代码；运行程序并排错。','改了哪些文件？运行结果是什么？']],[0.23,0.39,0.38]),
 h('认识一个简化的工作过程'),
 note('接收目标 → 读取上下文 → 选择下一步 → 调用可用工具 → 观察结果 → 调整或结束'),
 p('工具可以是文件读取、代码编辑、终端命令或其他被允许的操作。不同工具和模式可执行的动作不同，以界面实际提供的能力为准。'),
 h('课堂观察三个层次'),
 p('回答：“可以增加借用功能。”\n计划：“准备增加申请页面和审批逻辑。”\n执行证据：实际改动、启动输出、输入与返回结果。'),
 p('看到“已完成”后，继续检查：有没有运行？是否符合需求？异常情况有没有处理？未实际验证的内容，记录为“未验证”。'),
 h('本课中的 Agent 扮演什么角色'),
 p('本课先把 Agent 当作开发助手，帮助我们搭建普通仓库系统。给仓库系统加入自然语言查询，是以后可以探索的功能。'),
 note('想一想：Agent 生成了一张漂亮页面，是否就说明借还流程已经可用？你会怎样验证？'),
),
page('Prompt 与上下文：把任务说清楚', '少一点猜测，多一点可检查的信息',
 h('从一句简单要求开始'),
 prompt('帮我做一个仓库管理系统。'),
 p('先观察缺少什么：谁使用、解决什么问题、已有材料、第一版范围、运行环境，以及怎样判断完成。'),
 h('补充信息后的需求提示词'),
 prompt('背景：我们准备做一个学校仓库系统，学生刚学过 Java 基础。\n使用者：借用人和仓库管理员，具体权限待讨论。\n候选功能：物品管理、登录、借还审批、库存调整、统计报表。\n任务：先问我最多 5 个关键问题，再整理简版需求卡。\n要求：区分已知事实、你的假设和待确认项；给出第一版范围建议与验收标准。这一轮先不生成代码。'),
 h('提示词里放什么'),
 table(['要素','可以提供的内容'],[
 ['目标','本轮只要完成哪一件事？'],
 ['上下文','项目背景、已确认规则、现有代码、环境与错误原文。'],
 ['约束','沿用什么结构？本轮修改哪些内容？哪些暂不做？'],
 ['输出与验收','需要需求表、方案还是代码？用哪些操作检查结果？']],[0.24,0.76]),
 h('上下文要跟着任务更新'),
 p('第一轮讨论需求；下一轮带上已确认的需求与工程状态；出错时补充输入、预期、实际结果和错误信息。新开对话时重新提供关键背景，不默认工具记得所有决定。'),
 note('练习：给“加一个借用功能”补充用户、业务规则和一条验收条件。模板文件里的方括号需要替换成真实情况。'),
),
page('模拟甲乙方：把想法变成需求', '先问清楚，再评估和开发',
 prompt('甲方：我想知道仓库有什么东西、谁借走了、什么时候归还，还希望能查看统计。'),
 p('甲方提出场景、规则和优先级；乙方追问、复述理解、评估实现方式；双方一起确认范围和验收。课堂轮换角色，每个人都要提出一个问题。'),
 h('乙方可以从这些问题开始'),
 p('谁可以查询、借用、审批和调整库存？\n审批通过就算借出，还是实际领取后才算？\n什么时候减少可用数量？归还由谁确认？\n库存不足、重复提交、取消申请时怎么办？\n第一版最重要的一个流程是什么？'),
 h('项目过程与本课产出'),
 table(['环节','你们要留下什么'],[
 ['需求分析','用户、场景、业务规则与验收条件。'],
 ['范围评估','本次必做项、以后再做项。'],
 ['可行性与工作量评估','已会的、需要学习的、依赖与不确定性。'],
 ['方案设计','页面、Java 逻辑和数据库的职责草图。'],
 ['任务拆分','小任务、负责人、前置条件与粗略估时。'],
 ['开发与验证','实际改动、正常与异常测试记录。'],
 ['验收与交付','演示、启动方法、已完成与已知问题。']],[0.34,0.66]),
 note('不要把 Agent 的假设直接当作甲方要求。把未确定的问题写下来，课堂确认后再进入开发。以上是教学用简化过程，可以随反馈调整。'),
),
page('仓库功能：从整体范围到小步实现', '先讨论全貌，再选择第一轮任务',
 table(['候选模块','需要澄清的问题'],[
 ['物品基础管理','有哪些字段？如何分类、查询和维护？'],
 ['登录与角色权限','谁能做什么？不同角色怎样区分？'],
 ['借还审批','申请、审批、领取、归还各是什么状态？'],
 ['库存调整','何时增减？谁可调整？是否记录原因和操作者？'],
 ['统计报表','统计库存、借用还是逾期？时间范围和口径是什么？']],[0.32,0.68]),
 h('下面是讨论用数据，不是已经确定的表结构'),
 table(['编号','物品','位置','可用数量'],[
 ['A001','键盘','A-01','5'],['A002','螺丝刀','B-02','12'],['A003','安全帽','C-01','8']],[0.20,0.30,0.25,0.25]),
 p('例如：“库存数量”是总量、在库数量还是可借数量？借出 2 件以后页面显示什么？先说清定义，再决定需要保存哪些数据。'),
 h('第一轮可以怎样缩小'),
 p('一个起点示例：启动系统 → 查看物品列表 → 输入关键词 → 显示结果。课堂先由甲乙方确认本轮范围，再决定是否采用这个起点。'),
 p('随后可以逐步增加登录与权限、借还流程、库存调整和统计报表。每一轮都记录：新增了什么、验证了什么、还有什么问题。'),
 h('什么叫“完成”'),
 p('把“搜索好用”改写成可检查的行为。例如：数据中存在“键盘”，输入“键”后应显示它；输入空白时如何提示、找不到时如何反馈，也需要约定。'),
 note('今天不以功能数量评判结果。把一项需求说清、实现并验证，或者准确定位阻塞，都应留下可检查的过程记录。'),
),
page('遇到问题时，把知识接起来', '本页用于认识联系，不要求课前配置整套环境',
 table(['遇到的问题','将用到的知识'],[
 ['物品需要名称、位置、数量','Java 类、对象与集合。'],
 ['借用规则或状态处理不正确','方法、条件判断与业务规则。'],
 ['希望在浏览器里使用系统','JavaWeb、HTTP 请求与响应、接口。'],
 ['页面需要读取查询结果','JSON 数据格式与前后端交互。'],
 ['数据希望在程序重启后保留','MySQL、表结构与 Java 数据访问。'],
 ['Agent 改完后运行失败','错误日志、代码改动检查与逐步排查。']],[0.49,0.51]),
 h('先认识这条数据路线'),
 note('浏览器提交查询 → Java 接收并处理 → 读取数据 → Java 返回结果 → 页面展示\n读取的数据可以先来自示例列表，之后再接入 MySQL。'),
 h('课前环境检查'),
 p('□ 找到上一课项目，尝试运行、搜索与退出。\n□ 能指出读取输入、方法调用和列表遍历的位置。\n□ 能打开一个自己已有的 AI/Agent 工具；若不能，记录原因。\n□ 把当前代码和报错带来，先保留可复现的状态。'),
 p('继续使用上一课已经配置好的 Java 环境。本次课前不要求自行安装 MySQL、搭建 Spring Boot 工程、购买工具或申请 API Key；暂未配置 Agent 的同学可课堂结对操作。'),
 h('带来三样东西'),
 p('1. 工作单第 1 页：一个真实场景、三个待确认问题、一项优先功能。\n2. 自己的旧程序与运行状态：能运行或具体卡点。\n3. 一个想在课堂上弄明白的问题。'),
 note('工作单第 2–4 页用于课堂，不需要提前填写。你可以直接编辑 Word，也可以在纸上写下第 1 页的答案。'),
)
]

WORKBOOK = [
page('课前需求思考单', '第 1 页课前填写 · 建议 5–10 分钟 · 不确定处写“待确认”',
 p('姓名：________________    日期：________________    小组：________________'),
 h('1 / 我遇到的仓库使用场景'),
 p('可以选择学校器材、实验室工具或其他熟悉场景。没有经历也可以写一个假设场景，并标明“假设”。'),
 lines(['谁在什么情况下遇到什么问题？','目前怎样处理？哪里不方便？']),
 h('2 / 我会向甲方提出的三个问题'),
 lines(['问题 1：','问题 2：','问题 3：']),
 h('3 / 我认为第一版最重要的一项功能'),
 lines(['功能与选择理由：','怎样操作、看到什么结果，才算完成？']),
 h('4 / 我的准备状态'),
 p('上一课程序：□ 能运行  □ 有报错  □ 找不到项目\n已有 AI/Agent 工具：________________  □ 能使用  □ 尚未配置\n如有卡点，记录具体错误或现场带来：____________________________'),
 note('课前只需完成这一页。无需提前生成项目，也不要求预先确定所有业务规则。'),
),
page('甲乙方需求确认卡', '课堂填写 · 先标明事实、假设与待确认项',
 p('小组：____________  甲方代表：____________  乙方代表：____________'),
 lines(['用户与使用场景：','当前最需要解决的问题：']),
 table(['候选功能','本次/以后/待定','需要确认的规则'],[
 ['物品基础管理','',''],['登录与角色权限','',''],['借还审批','',''],['库存调整','',''],['统计报表','','']],[0.30,0.28,0.42]),
 h('关键规则'),
 lines(['谁能审批和调整库存？','何时算实际借出？何时增减可用数量？','库存不足、取消或重复提交时怎么办？']),
 h('信息状态'),
 lines(['已确认：','暂时假设：','待确认的问题与负责人：']),
 note('双方复述：本次为谁解决什么问题？哪些规则仍然不确定？未确认项不能悄悄变成已确认需求。'),
),
page('第一轮范围、评估与验收', '课堂填写 · 一轮只选择可以运行和检查的小任务',
 lines(['本轮目标：','必做项：','本轮暂不做：']),
 table(['评估项','本组判断'],[
 ['技术与能力','已会：                         需要学习：'],
 ['环境与数据','已具备：                       还缺：'],
 ['不确定性','可能的阻塞：                   应对：']],[0.28,0.72]),
 h('任务拆分与估时（估计，不作为完成承诺）'),
 table(['任务','负责人','前置条件','粗估分钟/不确定'],[
 ['','','',''],['','','',''],['','','','']],[0.33,0.16,0.24,0.27]),
 h('先写验收，再让 Agent 实现'),
 table(['场景','前提与操作','预期结果'],[
 ['正常情况','',''],['异常或边界','',''],['另一项检查','','']],[0.23,0.40,0.37]),
 lines(['甲方确认的范围与仍待确认事项：']),
),
page('Agent 搭建与问题记录', '课堂填写 · 需要下一轮时复制此页',
 p('轮次：________  本人负责：________________  使用工具：________________'),
 lines(['本轮给 Agent 的目标与约束：','提供了哪些上下文、文件或错误信息？','实际修改或生成了什么？']),
 h('运行与验证'),
 table(['输入/操作','预期结果','实际结果/未运行','结论'],[
 ['','','',''],['','','',''],['','','','']],[0.25,0.25,0.32,0.18]),
 p('状态：□ 已运行并验证  □ 能启动但结果不对  □ 未能启动  □ 尚未运行'),
 h('留下一个可继续排查的问题'),
 lines(['错误原文或具体现象：','已尝试的操作：','下一步准备查什么？需要补哪项知识？']),
 h('离场前，用自己的话说明'),
 lines(['我本人完成并能解释的是：','当前已完成/尚未完成：']),
 note('即使没跑通，也要记录启动步骤、具体阻塞和下一步。仅有 Agent 的“已完成”回复，不计作运行证据。'),
)
]

PROMPTS = '''第 2 课｜Agent 提示词模板
方括号请替换为真实情况；它们是模板，不是已确定需求。

一、先澄清需求
背景：我们要为[使用者]做一个仓库系统，当前问题是[场景与痛点]。
已知材料：[已有规则、原始需求或资料]。
候选功能：物品基础管理、登录、借还审批、库存调整、统计报表。
请先问最多 5 个关键问题。根据我的回答区分已知事实、假设和待确认项，整理简版需求卡，并建议第一版范围与验收标准。先不要生成代码。

二、准备第一轮搭建
已确认需求：[粘贴确认后的需求卡]。
当前工程与环境：[项目位置、目录结构、Java 环境、启动方式；未知处写未知]。
本轮目标：[只写一项可验证的小任务]。
约束：[沿用的技术/结构、允许修改范围、暂不做项]。
验收：[具体输入、操作、预期结果]。
先检查已有文件和环境，说明最小实现计划；缺少关键信息时先指出。实施后列出改动、运行步骤和实际验证结果。示例数据请明确标注，未运行的检查标记为未验证。

三、在已有系统上增加功能
已有能力：[已验证功能]。
新增目标：[一项功能]。
业务规则：[角色、状态变化、数据变化、异常处理]。
保持：[仍需正常工作的已有行为]。
请先检查相关代码，再小步修改，并验证新增功能及受影响的原有功能。记录未解决问题，不以“已完成”代替验证结果。

四、排查一个具体问题
我正在完成：[任务]。
启动/操作步骤：[可复现步骤]。
输入：[输入内容]。
预期：[本来应该发生什么]。
实际：[实际发生什么]。
错误原文与相关代码：[粘贴内容]。
已经试过：[操作与结果]。
请先分析可能原因和定位依据，再提出局部修复。修复后重做原来的复现步骤，写明实际结果。

五、整理本轮交付
请根据本轮真实改动和运行记录，整理已实现、已验证、未验证、已知问题、启动方法和下一步建议。不要把假设或未运行的功能写成已完成。
'''

STYLES = {
 'body': ParagraphStyle('body',fontName=FONT,fontSize=10.6,leading=17.8,textColor=HexColor(INK),wordWrap='CJK',spaceAfter=9),
 'h': ParagraphStyle('h',fontName=FONT,fontSize=13,leading=19,textColor=HexColor(TEAL),spaceBefore=10,spaceAfter=7),
 'cell': ParagraphStyle('cell',fontName=FONT,fontSize=9.5,leading=15.5,textColor=HexColor(INK),wordWrap='CJK'),
 'th': ParagraphStyle('th',fontName=FONT,fontSize=9.5,leading=15,textColor=white,wordWrap='CJK'),
 'prompt': ParagraphStyle('prompt',fontName=FONT,fontSize=10,leading=17,textColor=HexColor(INK),wordWrap='CJK'),
 'title': ParagraphStyle('title',fontName=FONT,fontSize=22,leading=29,textColor=HexColor(INK),spaceAfter=8),
 'sub': ParagraphStyle('sub',fontName=FONT,fontSize=10,leading=16,textColor=HexColor(MUTED),spaceAfter=16),
}
def para(text, kind='body'):
 return Paragraph(escape(text).replace('\n','<br/>'),STYLES[kind])

class PageCanvas(canvas.Canvas):
 def __init__(self,*args,**kwargs):
  self.course_label = kwargs.pop('course_label','')
  super().__init__(*args,**kwargs)
 def showPage(self):
  w, height = A4
  self.setStrokeColor(HexColor(TEAL)); self.setLineWidth(2)
  self.line(44,height-39,w-44,height-39)
  self.setFont(FONT,8); self.setFillColor(HexColor(MUTED))
  self.drawString(44,height-29,'103 仓库实战  /  LESSON 02')
  self.drawRightString(w-44,height-29,self.course_label)
  self.setStrokeColor(HexColor('#DDE5E9')); self.setLineWidth(.5)
  self.line(44,41,w-44,41)
  self.drawString(44,27,'学生资料 · 2026.10.20 · v1.0')
  self.drawRightString(w-44,27,f'{self._pageNumber:02d}')
  super().showPage()

def make_pdf(pages, path, label):
 compact = pages is WORKBOOK
 if compact:
  STYLES['body'].spaceAfter=6; STYLES['body'].leading=16.5
  STYLES['h'].spaceBefore=8; STYLES['h'].spaceAfter=5
  STYLES['sub'].spaceAfter=12
 doc = SimpleDocTemplate(str(path),pagesize=A4,rightMargin=44,leftMargin=44,topMargin=58,bottomMargin=56,title=label,author='103 仓库课程',allowSplitting=0)
 flow=[]; width=A4[0]-88
 for i,pg in enumerate(pages):
  if i: flow.append(PageBreak())
  flow += [para(pg['title'],'title'),para(pg['subtitle'],'sub')]
  for b in pg['blocks']:
   typ=b['type']
   if typ in ['p','h']: flow.append(para(b['text'],'body' if typ=='p' else 'h'))
   elif typ in ['note','prompt']:
    box=Table([[para(b['text'],'prompt')]],colWidths=[width])
    box.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),HexColor(PALE)),('BOX',(0,0),(-1,-1),.5,HexColor('#D9E8EA')),('LEFTPADDING',(0,0),(-1,-1),12),('RIGHTPADDING',(0,0),(-1,-1),12),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]))
    flow.extend([box,Spacer(1,10)])
   elif typ=='table':
    rows=[[para(x,'th') for x in b['headers']]]+[[para(x,'cell') for x in row] for row in b['rows']]
    rowheights=[None]+[30 if all(x=='' for x in row) else None for row in b['rows']]
    t=Table(rows,colWidths=[width*x for x in b['widths']],rowHeights=rowheights,repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor(TEAL)),('ROWBACKGROUNDS',(0,1),(-1,-1),[white,HexColor('#F3F6F8')]),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.4,HexColor('#D8E2E8')),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    flow.extend([t,Spacer(1,10)])
   elif typ=='lines':
    for text in b['labels']:
     flow.append(para(text))
     line=Table([['']],colWidths=[width],rowHeights=[10 if compact else 12])
     line.setStyle(TableStyle([('LINEBELOW',(0,0),(-1,-1),.4,HexColor('#B8C6D2'))]))
     flow.extend([line,Spacer(1,6 if compact else 8)])
 doc.build(flow,canvasmaker=lambda *a,**kw:PageCanvas(*a,course_label=label,**kw))
 actual=len(PdfReader(path).pages)
 if actual!=len(pages): raise RuntimeError(f'{path.name}: expected {len(pages)} pages, got {actual}')

def font(run, size=10.5, color=INK, bold=False):
 run.font.name='Microsoft YaHei'; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
 run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

def shade(cell, fill):
 el=OxmlElement('w:shd');el.set(qn('w:fill'),fill.lstrip('#'));cell._tc.get_or_add_tcPr().append(el)

def make_docx(pages,path,label):
 doc=Document(); sec=doc.sections[0]
 sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.65);sec.right_margin=Cm(1.65)
 normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10)
 normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(5)
 hp=sec.header.paragraphs[0];font(hp.add_run('103 仓库实战  /  LESSON 02'),8,MUTED)
 fp=sec.footer.paragraphs[0];font(fp.add_run('学生资料 · 2026.10.20 · v1.0     |     '),8,MUTED)
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
 doc.core_properties.title=label;doc.core_properties.author='103 仓库课程'
 for i,pg in enumerate(pages):
  if i: doc.add_page_break()
  q=doc.add_paragraph();q.paragraph_format.space_after=Pt(5);font(q.add_run(pg['title']),19,INK,True)
  q=doc.add_paragraph();font(q.add_run(pg['subtitle']),9,MUTED)
  for b in pg['blocks']:
   typ=b['type']
   if typ in ['p','h','note','prompt']:
    q=doc.add_paragraph();q.paragraph_format.keep_together=True
    if typ=='h': q.paragraph_format.keep_with_next=True; q.paragraph_format.space_before=Pt(7)
    if typ in ['note','prompt']:
     pr=q._p.get_or_add_pPr();el=OxmlElement('w:shd');el.set(qn('w:fill'),'EFF6F7');pr.append(el)
    font(q.add_run(b['text']),12 if typ=='h' else 10,TEAL if typ=='h' else INK,typ=='h')
   elif typ=='table':
    t=doc.add_table(rows=1,cols=len(b['headers']));t.style='Table Grid';t.autofit=False
    for col,ratio in zip(t.columns,b['widths']): col.width=Cm(17.7*ratio)
    for j,text in enumerate(b['headers']):
     c=t.rows[0].cells[j];shade(c,TEAL);font(c.paragraphs[0].add_run(text),9.2,'#FFFFFF',True)
    for row in b['rows']:
     cells=t.add_row().cells
     for j,text in enumerate(row):
      font(cells[j].paragraphs[0].add_run(text or ' '),9.2)
      cells[j].paragraphs[0].paragraph_format.space_after=Pt(3)
     if all(x=='' for x in row): t.rows[-1].height=Cm(.85)
    for row in t.rows:
     for cell, ratio in zip(row.cells,b['widths']): cell.width=Cm(17.7*ratio)
     pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
   elif typ=='lines':
    for text in b['labels']:
     q=doc.add_paragraph();q.paragraph_format.keep_with_next=True;font(q.add_run(text),10)
     q=doc.add_paragraph();font(q.add_run('__________________________________________________________________'),9,MUTED)
     q.paragraph_format.space_after=Pt(5)
 doc.save(path)

def to_markdown(pages):
 out=[]
 for i,pg in enumerate(pages,1):
  out.extend([f'# {i}. {pg["title"]}',f'> {pg["subtitle"]}',''])
  for b in pg['blocks']:
   typ=b['type']
   if typ=='p':out.extend([b['text'],''])
   elif typ=='h':out.extend(['## '+b['text'],''])
   elif typ=='note':out.extend(['> '+b['text'].replace('\n','\n> '),''])
   elif typ=='prompt':out.extend(['```text',b['text'],'```',''])
   elif typ=='table':
    out.append('| '+' | '.join(b['headers'])+' |');out.append('| '+' | '.join(['---']*len(b['headers']))+' |')
    out.extend('| '+' | '.join(row)+' |' for row in b['rows']);out.append('')
   elif typ=='lines':out.extend(text+'\n\n__________________________________________________\n' for text in b['labels'])
  out.append('---\n')
 return '\n'.join(out)

def main():
 for d in [STUDENT,SOURCE,QA]:d.mkdir(exist_ok=True)
 (SOURCE/'materials.json').write_text(json.dumps({'handbook':HANDBOOK,'workbook':WORKBOOK},ensure_ascii=False,indent=2),encoding='utf-8')
 (SOURCE/'01-课前预习手册.md').write_text(to_markdown(HANDBOOK),encoding='utf-8')
 (SOURCE/'02-仓库需求思考与课堂记录.md').write_text(to_markdown(WORKBOOK),encoding='utf-8')
 (STUDENT/'03-Agent提示词模板.txt').write_text(PROMPTS,encoding='utf-8-sig')
 make_pdf(HANDBOOK,STUDENT/'01-课前预习手册.pdf','课前预习手册')
 make_docx(HANDBOOK,SOURCE/'01-课前预习手册-教师可编辑版.docx','课前预习手册')
 make_pdf(WORKBOOK,STUDENT/'02-仓库需求思考与课堂记录.pdf','需求思考与课堂记录')
 make_docx(WORKBOOK,STUDENT/'02-仓库需求思考与课堂记录.docx','需求思考与课堂记录')
 previews=[]
 for filename in ['01-课前预习手册.pdf','02-仓库需求思考与课堂记录.pdf']:
  pdf=pdfium.PdfDocument(STUDENT/filename)
  for i in range(len(pdf)):
   image=pdf[i].render(scale=1.0).to_pil().convert('RGB')
   image.save(QA/f'{filename[:2]}-page-{i+1:02d}.png')
   thumb=ImageOps.contain(image,(298,422))
   card=Image.new('RGB',(318,456),'#e6ecef');card.paste(thumb,((318-thumb.width)//2,8))
   ImageDraw.Draw(card).text((12,435),f'{filename[:2]} / {i+1}',fill='#203348')
   previews.append(card)
  pdf.close()
 cols=3;rows=(len(previews)+cols-1)//cols
 contact=Image.new('RGB',(cols*318,rows*456),'#d4dfe5')
 for i,img in enumerate(previews):contact.paste(img,((i%cols)*318,(i//cols)*456))
 contact.save(QA/'contact-sheet.png')
 print('Generated: handbook 6 pages; workbook 4 pages; editable Word; prompts; source; page previews.')

if __name__=='__main__':main()


