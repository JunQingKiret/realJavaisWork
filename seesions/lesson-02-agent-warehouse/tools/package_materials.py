from pathlib import Path
import hashlib, json, re, zipfile
from pypdf import PdfReader
from docx import Document
from lxml import etree
import pdfplumber

ROOT=Path(__file__).resolve().parents[1]
STUDENT=ROOT/'student'
data=json.loads((ROOT/'source/materials.json').read_text(encoding='utf-8'))
clean=lambda text:re.sub(r'\s+','',text)

def snippets(page):
 yield page['title'];yield page['subtitle']
 for block in page['blocks']:
  if 'text' in block:yield block['text']
  if block['type']=='table':
   yield from block['headers']
   for row in block['rows']:yield from row
  if block['type']=='lines':yield from block['labels']

checks=[]
for filename,key in [('01-课前预习手册.pdf','handbook'),('02-仓库需求思考与课堂记录.pdf','workbook')]:
 path=STUDENT/filename
 reader=PdfReader(path)
 assert len(reader.pages)==len(data[key])
 for i,(actual,expected) in enumerate(zip(reader.pages,data[key]),1):
  text=clean(actual.extract_text())
  for snippet in snippets(expected):
   assert clean(snippet) in text,(filename,i,snippet)
  assert '\ufffd' not in text
 with pdfplumber.open(path) as pdf:
  for page in pdf.pages:
   assert all(c['x0']>=20 and c['x1']<=page.width-20 and c['top']>=10 and c['bottom']<=page.height-10 for c in page.chars),filename
 checks.append(f'{filename}：{len(reader.pages)} 页，逐页文字与源内容一致，字符未越出页面边界。')

for path,key in [(STUDENT/'02-仓库需求思考与课堂记录.docx','workbook'),(ROOT/'source/01-课前预习手册-教师可编辑版.docx','handbook')]:
 doc=Document(path)
 text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
 normalized=clean(text)
 for page in data[key]:
  for snippet in snippets(page):assert clean(snippet) in normalized,(path.name,snippet)
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None
  root=etree.fromstring(z.read('word/document.xml'))
  breaks=root.xpath('count(//w:br[@w:type="page"])',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
  assert int(breaks)==len(data[key])-1
 checks.append(f'{path.name}：文本、表格、显式分页结构与源内容一致，DOCX 文件结构检查通过。')

expected={'00-阅读顺序.txt','01-课前预习手册.pdf','02-仓库需求思考与课堂记录.docx','02-仓库需求思考与课堂记录.pdf','03-Agent提示词模板.txt'}
files=list(STUDENT.iterdir())
assert {p.name for p in files}==expected
assert all(p.is_file() and p.stat().st_size>0 for p in files)
for path in files:
 if path.suffix=='.txt':
  text=path.read_text(encoding='utf-8-sig')
 elif path.suffix=='.pdf':
  text='\n'.join(p.extract_text() for p in PdfReader(path).pages)
 else:
  doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
 for forbidden in ['E:\\','C:\\Users\\','codex://','TODO','教师答案']:
  assert forbidden not in text,(path.name,forbidden)

archive=ROOT/'lesson-02-agent-warehouse-student.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
 for path in sorted(files):z.write(path,'student/'+path.name)
with zipfile.ZipFile(archive) as z:
 assert z.testzip() is None
 assert set(z.namelist())=={'student/'+name for name in expected}
 for path in files:assert z.read('student/'+path.name)==path.read_bytes()
checks.append('学生压缩包仅含 student 目录中的 5 个预期文件；逐文件内容与原件一致，ZIP 完整性检查通过。')
checks.append('学生文件无本机用户路径、历史聊天链接、教师答案或遗留 TODO。')
checks.append('已目视检查 10 页 PDF 缩略图，以及课前工作单完整单页，未发现文字遮挡或异常分页。')
report=['# 第二节课前资料交付检查','', '检查日期：2026-10-20。','', '## 内容与格式检查','']
report += ['- '+x for x in checks]
report += ['', '## 检查边界','', '- Word 已做内容和结构检查，未用 Word/WPS 实际渲染；编辑器和字体差异可能引起自动分页。固定分页打印请使用 PDF。', '- 本次只制作课前学生材料，没有创建、构建或测试仓库应用。', '- 未代为发送班群消息。', '', '## 学生包清单','', '| 文件 | 字节数 |','|---|---|']
report += [f'| {p.name} | {p.stat().st_size} |' for p in sorted(files)]
report += ['',f'压缩包：`{archive.name}`',f'大小：{archive.stat().st_size} 字节',f'SHA-256：`{hashlib.sha256(archive.read_bytes()).hexdigest()}`','']
(ROOT/'交付检查记录.md').write_text('\n'.join(report),encoding='utf-8')
print('\n'.join(checks))
print(f'ZIP: {archive.stat().st_size} bytes; 5 files')
